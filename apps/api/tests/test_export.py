"""Тесты экспорта каталога в CSV/XLSX. См. ARCHITECTURE_PLAN.md §16 п.16.

Слои:
  1. Эндпоинт POST /catalog/export: 202 (csv/xlsx) + job QUEUED + dispatch
     Celery-задачи; 422 для pdf.
  2. Задача _run_export (БД реальная, S3 замокан): CSV c персональной ценой
     (utf-8-sig, разделитель «;»), XLSX читается openpyxl, статусы
     QUEUED → RUNNING → DONE; StorageError → FAILED (не роняет worker).
  3. GET /catalog/export/{job_id}: DONE + presigned url; чужой пользователь → 404;
     случайный job_id → 404.
  4. Фильтры: brand_ids ограничивает выгрузку товарами бренда.
"""
import asyncio
import csv
import io
import json
import uuid

import pytest
from openpyxl import load_workbook

from app.api.v1 import catalog as catalog_api
from app.core.config import settings
from app.models.enums import UserRole
from app.repositories.catalog import CatalogFilters
from app.services import export as export_service
from app.services.storage import StorageError
from app.tasks import export_catalog as export_task
from tests.conftest import (
    create_brand,
    create_product,
    create_series,
    create_user,
    set_discount,
)

PASSWORD = "Passw0rd!"
CLIENT_EMAIL = "client@example.by"
OTHER_EMAIL = "other@example.by"

COLUMNS = ["sku", "name", "brand", "series", "stock_status",
           "base_price", "unit_price", "currency"]


# ---------- helpers ----------

class FakeRedis:
    """In-memory Redis для job-стейта (используются только get/set с ex)."""

    def __init__(self):
        self.values = {}

    async def get(self, key):
        return self.values.get(key)

    async def set(self, key, value, ex=None):
        self.values[key] = value
        return True


@pytest.fixture(autouse=True)
def _fake_job_redis(monkeypatch):
    """Весь модуль — на FakeRedis: модульный redis-клиент сервиса не должен
    открывать реальное соединение (незакрытое asyncio-соединение рождает
    «Event loop is closed» при выходе из pytest)."""
    monkeypatch.setattr(export_service, "_redis", FakeRedis())


@pytest.fixture
def job_redis(monkeypatch):
    """Подменяем redis-клиент сервиса экспорта (паттерн _isolated_cache)."""
    fake = FakeRedis()
    monkeypatch.setattr(export_service, "_redis", fake)
    return fake


@pytest.fixture
def dispatched(monkeypatch):
    """Перехват dispatch Celery-задачи (broker в тестах недоступен)."""
    calls: list[tuple] = []
    monkeypatch.setattr(
        export_service.run_export, "delay", lambda *a, **k: calls.append(a)
    )
    return calls


def _capture_upload(monkeypatch):
    """Мокаем S3: upload_fileobj складывает тело файла в список."""
    uploads = []

    def _upload(bucket, key, fileobj, **kwargs):
        uploads.append(
            {"bucket": bucket, "key": key, "body": fileobj.read(), "kwargs": kwargs}
        )

    monkeypatch.setattr(export_task.storage, "upload_fileobj", _upload)
    return uploads


async def _login(api_client, email):
    r = await api_client.post(
        "/api/v1/auth/login", json={"email": email, "password": PASSWORD}
    )
    assert r.status_code == 200, r.text


async def _seed_catalog(sf):
    """2 бренда, 2 товара; клиенту — скидка 10% на бренд KEAZ (unit_price 90)."""
    client = await create_user(sf, email=CLIENT_EMAIL, role=UserRole.CLIENT)
    keaz = await create_brand(sf, name="KEAZ")
    iek = await create_brand(sf, name="IEK")
    series = await create_series(sf, brand=keaz, name="Optibox")
    await create_product(sf, sku="A-1", name="Widget", brand=keaz, series=series,
                         base_price=100)
    await create_product(sf, sku="B-2", name="Gadget", brand=iek, base_price=200)
    await set_discount(sf, user=client, brand=keaz, percent=10)
    return client, keaz, iek


async def _run_task(sf, monkeypatch, *, job_id, user, filters, fmt):
    """Прогнать async-пайплайн задачи (S3 замокан, сессия — тестовый движок)."""
    monkeypatch.setattr(export_task, "_worker_session", sf)
    uploads = _capture_upload(monkeypatch)
    result = await export_task._run_export(
        uuid.UUID(job_id), user.id, filters, fmt, "fixed"
    )
    return result, uploads


def _job_state(job_redis, job_id) -> dict:
    return json.loads(job_redis.values[f"export:job:{job_id}"])


# ===========================================================================
# 1. POST /catalog/export
# ===========================================================================
@pytest.mark.asyncio
class TestStartEndpoint:
    async def test_start_csv_and_xlsx_202(
        self, api_client, session_factory, job_redis, dispatched
    ):
        client, keaz, _iek = await _seed_catalog(session_factory)
        await _login(api_client, CLIENT_EMAIL)

        for fmt in ("csv", "xlsx"):
            r = await api_client.post(
                "/api/v1/catalog/export",
                params={"format": fmt, "brand": str(keaz.id)},
            )
            assert r.status_code == 202, r.text
            job_id = r.json()["job_id"]
            assert uuid.UUID(job_id)  # валидный uuid

            # job создан в статусе QUEUED с параметрами выгрузки
            state = _job_state(job_redis, job_id)
            assert state["status"] == "QUEUED"
            assert state["user_id"] == str(client.id)
            assert state["format"] == fmt
            assert state["s3_key"] is None

            # Celery-задача диспатчнулась с этим же job_id
            assert dispatched[-1][0] == job_id

        # аргументы dispatch: (job_id, user_id, filters_json, format, mode)
        assert len(dispatched) == 2
        assert dispatched[0][1] == str(client.id)
        filters_json = json.loads(dispatched[0][2])
        assert filters_json["brand_ids"] == [str(keaz.id)]
        assert dispatched[0][3] == "csv" and dispatched[0][4] == "fixed"

    async def test_pdf_422(self, api_client, session_factory, job_redis, dispatched):
        await _seed_catalog(session_factory)
        await _login(api_client, CLIENT_EMAIL)
        r = await api_client.post("/api/v1/catalog/export", params={"format": "pdf"})
        assert r.status_code == 422, r.text
        assert "PDF" in r.json()["detail"]
        assert dispatched == []  # задача не запускалась


# ===========================================================================
# 2. Задача экспорта (БД + замоканный S3)
# ===========================================================================
@pytest.mark.asyncio
class TestExportTask:
    async def test_csv_with_personal_price(
        self, session_factory, monkeypatch, job_redis, dispatched
    ):
        client, keaz, _iek = await _seed_catalog(session_factory)
        filters = {"q": None, "brand_ids": [], "series_ids": [], "stock": None}
        job_id = await export_service.start_export(
            user=client, filters=CatalogFilters(), format="csv",
            price_calc_mode="fixed",
        )
        assert _job_state(job_redis, job_id)["status"] == "QUEUED"

        result, uploads = await _run_task(
            session_factory, monkeypatch, job_id=job_id, user=client,
            filters=filters, fmt="csv",
        )

        assert result["status"] == "DONE"
        assert result["rows"] == 2
        # statuses: QUEUED → RUNNING → DONE
        assert _job_state(job_redis, job_id)["status"] == "DONE"
        assert _job_state(job_redis, job_id)["s3_key"] == f"export/{job_id}.csv"

        bucket, key = uploads[0]["bucket"], uploads[0]["key"]
        assert bucket == settings.s3_bucket_exports
        assert key == f"export/{job_id}.csv"
        assert uploads[0]["kwargs"]["content_type"] == "text/csv"

        body = uploads[0]["body"]
        assert body.startswith(b"\xef\xbb\xbf")  # BOM: Excel откроет кириллицу
        rows = list(csv.reader(io.StringIO(body.decode("utf-8-sig")), delimiter=";"))
        assert rows[0] == COLUMNS
        # сортировка по имени: Gadget < Widget
        assert rows[1] == ["B-2", "Gadget", "IEK", "",
                           "IN_STOCK", "200.0", "200.0", "BYN"]
        # A-1: скидка 10% на KEAZ → unit_price 90.0
        assert rows[2] == ["A-1", "Widget", "KEAZ", "Optibox",
                           "IN_STOCK", "100.0", "90.0", "BYN"]

    async def test_xlsx_readable_by_openpyxl(
        self, session_factory, monkeypatch, job_redis, dispatched
    ):
        client, _keaz, _iek = await _seed_catalog(session_factory)
        filters = {"q": None, "brand_ids": [], "series_ids": [], "stock": None}
        job_id = await export_service.start_export(
            user=client, filters=CatalogFilters(), format="xlsx",
            price_calc_mode="fixed",
        )

        result, uploads = await _run_task(
            session_factory, monkeypatch, job_id=job_id, user=client,
            filters=filters, fmt="xlsx",
        )

        assert result["status"] == "DONE"
        assert uploads[0]["key"] == f"export/{job_id}.xlsx"
        assert uploads[0]["kwargs"]["content_type"].startswith(
            "application/vnd.openxmlformats"
        )
        wb = load_workbook(io.BytesIO(uploads[0]["body"]))
        ws = wb.active
        assert [c.value for c in ws[1]] == COLUMNS
        assert ws.max_row == 3  # заголовок + 2 товара (сортировка по имени)
        assert ws["A2"].value == "B-2"
        assert float(ws["F2"].value) == 200.0
        assert float(ws["G2"].value) == 200.0
        assert ws["A3"].value == "A-1"
        assert float(ws["F3"].value) == 100.0
        assert float(ws["G3"].value) == 90.0  # персональная цена (скидка 10%)
        assert ws["H3"].value == "BYN"

    async def test_storage_error_marks_failed(
        self, session_factory, monkeypatch, job_redis, dispatched
    ):
        client, _keaz, _iek = await _seed_catalog(session_factory)
        filters = {"q": None, "brand_ids": [], "series_ids": [], "stock": None}
        job_id = await export_service.start_export(
            user=client, filters=CatalogFilters(), format="csv",
            price_calc_mode="fixed",
        )
        monkeypatch.setattr(export_task, "_worker_session", session_factory)

        def _boom(bucket, key, fileobj, **kwargs):
            raise StorageError("minio недоступен")

        monkeypatch.setattr(export_task.storage, "upload_fileobj", _boom)

        # sync-обёртка ловит ошибку → FAILED в Redis, worker не падает.
        # to_thread: обёртка зовёт asyncio.run — нельзя в loop текущего теста.
        res = await asyncio.to_thread(
            export_task.run_export,
            job_id, str(client.id), json.dumps(filters), "csv", "fixed",
        )
        assert res["status"] == "FAILED"
        state = _job_state(job_redis, job_id)
        assert state["status"] == "FAILED"
        assert "minio" in state["error"]

    async def test_filters_brand_only(
        self, session_factory, monkeypatch, job_redis, dispatched
    ):
        """Экспорт с brand_ids содержит только товары бренда."""
        client, keaz, _iek = await _seed_catalog(session_factory)
        filters = {"q": None, "brand_ids": [str(keaz.id)],
                   "series_ids": [], "stock": None}
        job_id = await export_service.start_export(
            user=client, filters=CatalogFilters(brand_ids=[keaz.id]),
            format="csv", price_calc_mode="fixed",
        )

        result, uploads = await _run_task(
            session_factory, monkeypatch, job_id=job_id, user=client,
            filters=filters, fmt="csv",
        )

        assert result["rows"] == 1
        text = uploads[0]["body"].decode("utf-8-sig")
        assert "A-1" in text and "KEAZ" in text
        assert "B-2" not in text and "IEK" not in text


# ===========================================================================
# 3. GET /catalog/export/{job_id}
# ===========================================================================
@pytest.mark.asyncio
class TestJobStatusEndpoint:
    async def test_done_returns_presigned_url(
        self, api_client, session_factory, monkeypatch, job_redis, dispatched
    ):
        client, _keaz, _iek = await _seed_catalog(session_factory)
        filters = {"q": None, "brand_ids": [], "series_ids": [], "stock": None}
        job_id = await export_service.start_export(
            user=client, filters=CatalogFilters(), format="csv",
            price_calc_mode="fixed",
        )
        await _run_task(
            session_factory, monkeypatch, job_id=job_id, user=client,
            filters=filters, fmt="csv",
        )
        await _login(api_client, CLIENT_EMAIL)
        monkeypatch.setattr(
            catalog_api.storage,
            "presigned_get",
            lambda bucket, key, **kw: f"http://minio.local/{key}",
        )

        r = await api_client.get(f"/api/v1/catalog/export/{job_id}")
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["job_id"] == job_id
        assert body["status"] == "DONE"
        assert body["format"] == "csv"
        assert body["url"] == f"http://minio.local/export/{job_id}.csv"

    async def test_other_user_404(
        self, api_client, session_factory, monkeypatch, job_redis, dispatched
    ):
        client, _keaz, _iek = await _seed_catalog(session_factory)
        job_id = await export_service.start_export(
            user=client, filters=CatalogFilters(), format="csv",
            price_calc_mode="fixed",
        )
        await create_user(session_factory, email=OTHER_EMAIL, role=UserRole.CLIENT)
        await _login(api_client, OTHER_EMAIL)

        r = await api_client.get(f"/api/v1/catalog/export/{job_id}")
        assert r.status_code == 404, r.text

    async def test_unknown_job_404(self, api_client, session_factory):
        await create_user(session_factory, email=CLIENT_EMAIL, role=UserRole.CLIENT)
        await _login(api_client, CLIENT_EMAIL)
        r = await api_client.get(f"/api/v1/catalog/export/{uuid.uuid4()}")
        assert r.status_code == 404, r.text
